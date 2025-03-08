import io
import PyPDF2
from docx import Document
import pandas as pd
from fpdf import FPDF
import streamlit as st

def extract_text_from_pdf(file):
    try:
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        st.error(f"Error extracting text from PDF: {str(e)}")
        return None

def extract_text_from_docx(file):
    try:
        doc = Document(file)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        st.error(f"Error extracting text from DOCX: {str(e)}")
        return None

def extract_text_from_file(uploaded_file):
    try:
        if uploaded_file.type == "application/pdf":
            return extract_text_from_pdf(uploaded_file)
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return extract_text_from_docx(uploaded_file)
        else:
            # For text files
            return uploaded_file.getvalue().decode("utf-8")
    except Exception as e:
        st.error(f"Error extracting text from file: {str(e)}")
        return None

def generate_pdf_report(analysis_results):
    pdf = FPDF()
    pdf.add_page()

    # Header
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Terms and Conditions - Agent Analysis Report", ln=True, align='C')
    pdf.line(10, 20, 200, 20)  # Add a horizontal line under header
    pdf.ln(5)

    # Analysis Metrics Summary
    if 'metrics' in analysis_results:
        metrics = analysis_results['metrics']
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(200, 10, txt="Analysis Summary", ln=True)
        pdf.set_font("Arial", size=10)
        pdf.cell(60, 8, txt=f"Complexity Score: {metrics['complexity_score']:.1f}%", border=1)
        pdf.cell(60, 8, txt=f"Financial Impact: {metrics['financial_impact']:.1f}%", border=1)
        pdf.cell(60, 8, txt=f"Unusual Terms: {metrics['unusual_terms_ratio']:.1f}%", border=1, ln=True)
        pdf.ln(5)

    # Section summaries with compact layout
    sections = {
        "Core Terms": ANALYSIS_CATEGORIES[:14],
        "Quality & Compliance": ANALYSIS_CATEGORIES[14:22],
        "Delivery & Fulfillment": ANALYSIS_CATEGORIES[22:]
    }

    current_section = None
    for category, details in analysis_results.items():
        if category == 'metrics' or category == 'metadata':
            continue

        # Determine section
        if category in ANALYSIS_CATEGORIES[:14]:
            section = "Core Terms"
        elif category in ANALYSIS_CATEGORIES[14:22]:
            section = "Quality & Compliance"
        else:
            section = "Delivery & Fulfillment"

        # Add section header if new section
        if section != current_section:
            current_section = section
            pdf.set_font("Arial", 'B', 11)
            pdf.cell(200, 8, txt=section, ln=True, fill=True)
            pdf.set_font("Arial", size=9)

        # Only include high and medium risk items
        if details['risk_level'] in ['High', 'Medium']:
            # Category name with risk level
            pdf.set_font("Arial", 'B', 10)
            risk_indicator = "[!]" if details['risk_level'] == "High" else "[*]"
            pdf.cell(200, 6, txt=f"{risk_indicator} {category}", ln=True)

            # Findings
            pdf.set_font("Arial", size=9)
            pdf.multi_cell(0, 4, txt=f"Findings: {details['findings']}")

            # Quoted terms
            if details['quoted_phrases']:
                terms = []
                for phrase in details['quoted_phrases']:
                    marker = "[F]" if phrase['is_financial'] else "[-]"
                    terms.append(f"{marker} {phrase['text']}")
                pdf.multi_cell(0, 4, txt="Terms: " + "\n".join(terms))

            pdf.ln(2)

    # Add legend at the bottom
    pdf.ln(5)
    pdf.set_font("Arial", 'I', 8)
    pdf.multi_cell(0, 4, txt="Legend: [!] High Risk  [*] Medium Risk  [F] Financial Term  [-] Non-Financial Term")

    return pdf.output(dest='S').encode('latin-1')

def generate_csv_report(analysis_results):
    # Add section info to each row
    data = []

    for category, details in analysis_results.items():
        if category == 'metrics':
            continue

        if category in ANALYSIS_CATEGORIES[:14]:
            section = "Core Terms"
            section_summary = "Fundamental agreement elements and basic rights"
        elif category in ANALYSIS_CATEGORIES[14:22]:
            section = "Quality & Compliance"
            section_summary = "Regulatory adherence and quality standards"
        else:
            section = "Delivery & Fulfillment"
            section_summary = "Logistics and delivery processes"

        # Create a row for each quoted phrase
        if details['quoted_phrases']:
            for phrase in details['quoted_phrases']:
                data.append({
                    'Section': section,
                    'Section Summary': section_summary,
                    'Category': category,
                    'Risk Level': details['risk_level'],
                    'Findings': details['findings'],
                    'Term': phrase['text'],
                    'Is Financial': 'Yes' if phrase['is_financial'] else 'No'
                })
        else:
            data.append({
                'Section': section,
                'Section Summary': section_summary,
                'Category': category,
                'Risk Level': details['risk_level'],
                'Findings': details['findings'],
                'Term': '',
                'Is Financial': ''
            })

    # Add metrics as a separate section if available
    if 'metrics' in analysis_results:
        metrics = analysis_results['metrics']
        data.append({
            'Section': 'Metrics Summary',
            'Section Summary': 'Overall analysis metrics',
            'Category': 'Complexity Score',
            'Risk Level': f"{metrics['complexity_score']:.1f}%",
            'Findings': 'Percentage of categories with specific requirements or unusual terms',
            'Term': '',
            'Is Financial': ''
        })

    df = pd.DataFrame(data)
    return df.to_csv(index=False).encode('utf-8')

ANALYSIS_CATEGORIES = ["Category1", "Category2", "Category3", "Category4", "Category5", "Category6", "Category7", "Category8", "Category9", "Category10", "Category11", "Category12", "Category13", "Category14", "Category15", "Category16", "Category17", "Category18", "Category19", "Category20", "Category21", "Category22", "Category23"]